"""
Fungsi bantu untuk idle lambda finder.
Ahmad Ahlul Hikam
"""
import re
import sys
import zipfile
import io
from pathlib import Path

import yaml
from openpyxl import load_workbook


def load_config(config_path: str) -> dict:
    path = Path(config_path)
    if not path.exists():
        print(f"[ERROR] Config tidak ditemukan: {config_path}")
        sys.exit(1)
    with open(path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def find_input_files(input_folder: str, pattern: str) -> list[Path]:
    folder = Path(input_folder)
    if not folder.exists():
        print(f"[ERROR] Folder input tidak ditemukan: {input_folder}")
        sys.exit(1)
    files = sorted(folder.glob(pattern))
    if not files:
        print(f"[ERROR] Tidak ada file '{pattern}' di {input_folder}")
        print("        Taruh file export 'Manage WDM Trail' (.xlsx) di folder itu.")
        sys.exit(1)
    return files


EMPTY_SHARED_STRINGS = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="0" uniqueCount="0"></sst>'
)


def open_workbook_safely(path: Path):
    """
    Buka file .xlsx dengan openpyxl. Beberapa export dari sistem NMS punya
    referensi ke 'xl/sharedStrings.xml' yang sebenarnya tidak ada di dalam
    file (bug dari sistem sumbernya) meski semua teks sudah disimpan inline.
    Ini menyebabkan openpyxl gagal dengan KeyError. Kalau itu terjadi,
    tambal file tersebut di memori dengan sharedStrings.xml kosong, baru
    dibuka lagi.
    """
    try:
        return load_workbook(path, data_only=True)
    except KeyError as e:
        if "sharedStrings" not in str(e):
            raise
        print(f"[INFO] Memperbaiki file rusak (sharedStrings hilang): {path.name}")
        with open(path, "rb") as f:
            original_bytes = f.read()

        buffer = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(original_bytes), "r") as zin:
            with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    zout.writestr(item, zin.read(item.filename))
                zout.writestr("xl/sharedStrings.xml", EMPTY_SHARED_STRINGS)

        buffer.seek(0)
        return load_workbook(buffer, data_only=True)


def find_header_row(ws, target_column: str, max_scan_rows: int = 30):
    """Cari baris header yang mengandung kolom target_column, kembalikan (row_idx, header_list)."""
    for row_idx in range(1, max_scan_rows + 1):
        row_values = [c.value for c in ws[row_idx]]
        if row_values and target_column in row_values:
            return row_idx, row_values
    return None, None


CHANNEL_PATTERN = re.compile(r"^[A-Za-z]?\\(\d+)\\")


def parse_channel_number(raw_value):
    """
    Ekstrak nomor channel dari format 'C\\74\\1558.17\\192.400' -> 74.
    Return None kalau formatnya tidak dikenali atau menandakan kosong ('-', 'C\\00\\...').
    """
    if raw_value is None:
        return None
    text = str(raw_value).strip()
    if text in ("-", "", "0", "00"):
        return None

    match = CHANNEL_PATTERN.match(text)
    if not match:
        return None

    channel_no = int(match.group(1))
    if channel_no == 0:
        return None
    return channel_no
